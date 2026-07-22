from argparse import Namespace
from dataclasses import replace
import tempfile
from unittest.mock import patch

import alr_quote_verifier as aqv


def _identity_link(**kwargs):
    candidate = kwargs.get("link_candidate") or ""
    if kwargs.get("kind") in {"case", "unreported", "statute"} and candidate.lower() == "other":
        return "https://example.test/resolved"
    return candidate


def _unresolved_link(**kwargs):
    return kwargs.get("link_candidate") or ""


def test_deterministic_pipeline_emits_full_reference_tuple():
    text = (
        "Johnson, supra note 243 at para 19 citing Lyons, supra note 243 at page 339; "
        "R v Standingwater, 2013 SKCA 78 at para 20."
    )
    with patch.object(aqv, "_resolve_footnote_part_link", side_effect=_identity_link):
        result = aqv._deterministic_footnote_parts(text)
    assert result is not None
    parts, reasons = result
    assert len(parts) == 3
    assert reasons == ("top_level_semicolon", "explicit_source_signal")
    assert parts[1].short_form == "Lyons"
    assert parts[1].page_pinpoints == [339]
    assert "supra note 243" in parts[1].bare_citation


def test_numbered_supra_from_deterministic_part_uses_existing_registry():
    with patch.object(aqv, "_resolve_footnote_part_link", side_effect=_identity_link):
        result = aqv._deterministic_footnote_parts(
            "Oakes, supra note 1 at para 20; R v Grant, 2009 SCC 32."
        )
    assert result is not None
    reference = [result[0][0]]
    registry = [{
        "verbatim": "R v Oakes, [1986] 1 SCR 103 [Oakes].",
        "link": "https://www.canlii.org/en/ca/scc/doc/1986/1986canlii46/1986canlii46.html#par1",
        "short_form": "Oakes",
        "note": "1",
    }]
    methods = aqv._resolve_footnote_reference_links(
        2, reference[0].verbatim, reference, registry,
        allow_fallback=False,
    )
    assert methods == {0: "note_number"}
    assert reference[0].link.endswith("#par20")


def test_ibid_from_deterministic_part_uses_sibling_authority():
    with patch.object(aqv, "_resolve_footnote_part_link", side_effect=_identity_link):
        result = aqv._deterministic_footnote_parts(
            "R v Oakes, [1986] 1 SCR 103; Ibid at para 25."
        )
    assert result is not None
    parts = result[0]
    parts[0] = replace(
        parts[0],
        link="https://www.canlii.org/en/ca/scc/doc/1986/1986canlii46/1986canlii46.html#par1",
    )
    methods = aqv._resolve_footnote_reference_links(
        1, "", parts, [], allow_fallback=False,
    )
    assert methods[1] == "sibling"
    assert parts[1].link.endswith("#par25")


def test_free_mode_is_lossless_when_splitter_abstains():
    text = "This explanatory note has no citation boundary."
    with patch.object(aqv, "_resolve_footnote_part_link", side_effect=_identity_link):
        result = aqv._deterministic_footnote_parts(text, allow_unsplit_fallback=True)
    assert result is not None
    assert [part.verbatim for part in result[0]] == [text]


def test_free_mode_uses_recall_first_semicolon_parts():
    text = "R v Oakes, [1986] 1 SCR 103; commentary; R v Grant, 2009 SCC 32."
    with patch.object(aqv, "_resolve_footnote_part_link", side_effect=_identity_link):
        result = aqv._deterministic_footnote_parts(text, allow_unsplit_fallback=True)
    assert result is not None
    assert [part.verbatim for part in result[0]] == [
        "R v Oakes, [1986] 1 SCR 103", "commentary", "R v Grant, 2009 SCC 32.",
    ]


def test_hrto_neutral_citation_builds_canlii_url():
    assert aqv._generate_fallback_url(
        "XY v Ontario (Government and Consumer Services), 2012 HRTO 726",
        kind="case",
    ) == (
        "https://www.canlii.org/en/on/onhrt/doc/2012/2012hrto726/"
        "2012hrto726.html"
    )


def test_a2aj_reporter_case_uses_authoritative_source_url():
    document = aqv.a2aj_client.A2AJDocument(
        dataset="SCC",
        citation="[1997] 3 SCR 484",
        alternate_citation="",
        name="R v S (RD)",
        date="1997-09-26",
        url="https://decisions.scc-csc.ca/scc-csc/scc-csc/en/item/1549/index.do",
        text="",
        language="en",
        scraped_timestamp="",
        upstream_license="",
        raw={},
    )
    assert aqv._a2aj_case_link(document, "en") == document.url


def test_a2aj_reporter_case_rejects_non_web_source_url():
    document = aqv.a2aj_client.A2AJDocument(
        dataset="SCC", citation="[1997] 3 SCR 484", alternate_citation="",
        name="R v S (RD)", date="1997-09-26", url="file:///tmp/case.html",
        text="", language="en", scraped_timestamp="", upstream_license="", raw={},
    )
    assert aqv._a2aj_case_link(document, "en") == ""


def test_journal_page_recovery_tolerates_whitespace_normalization():
    text = "[page 439]\nFirst page.\n[page 440]\nA quoted\npassage appears here."
    assert aqv._journal_db_page_for_region(
        text, "A quoted passage appears here."
    ) == "page 440"


def test_journal_page_recovery_reports_repeated_quote_pages():
    text = (
        "[page 439]\nRepeated phrase.\n"
        "[page 440]\nNo match here.\n"
        "[page 441]\nRepeated\nphrase."
    )
    assert aqv._journal_db_pages_for_quote(text, "Repeated phrase.") == [
        "page 439", "page 441",
    ]


def test_journal_page_recovery_uses_quote_match_threshold():
    text = "[page 449]\nThe source calls this a special responsibility."
    assert aqv._journal_db_pages_for_quote(
        text, "special responsibility,", min_score=0.85
    ) == ["page 449"]


def test_workbook_does_not_call_populated_pinpoint_missing():
    import openpyxl

    row = {
        "footnote_id": 1,
        "footnote_display_id": "1",
        "footnote_full": "Example footnote.",
        "proposition_text": 'A proposition with "quoted text".',
        "citation_part_text": "Example citation at para 20.",
        "has_quotes": "YES",
        "quote_check_status": "NO_MATCH",
        "quote_check_notes": "",
        "quote_corrected_citation": '"quoted text"',
        "citation_part_link": "https://example.test/source",
        "pinpoint_fragments": '["par20"]',
        "page_pinpoints": "[]",
    }
    with tempfile.TemporaryDirectory() as directory:
        path = f"{directory}\\result.xlsx"
        aqv.write_workbook({"footnote_rows": [row]}, path)
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet = workbook["FootnoteReferences"]
        headers = [cell.value for cell in sheet[1]]
        value = sheet.cell(2, headers.index("Corrected quote") + 1).value
        workbook.close()
    assert "No match found" in value
    assert "No pinpoint provided" not in value


def test_ultra_economy_leaves_mixed_reference_footnote_on_established_path():
    with patch.object(aqv, "_resolve_footnote_part_link", side_effect=_identity_link):
        result = aqv._deterministic_footnote_parts(
            "R v Oakes, [1986] 1 SCR 103; Ibid at para 25.",
            allow_reference_parts=False,
        )
    assert result is None


def test_ultra_requires_resolved_identity_but_free_does_not():
    text = "Dog Owners’ Liability Act, RSO 1990, c D.16, s 1 [DOLA]."
    with patch.object(aqv, "_resolve_footnote_part_link", side_effect=_unresolved_link):
        assert aqv._deterministic_footnote_parts(text) is None
        assert aqv._deterministic_footnote_parts(
            text, allow_unsplit_fallback=True
        ) is not None


def test_free_configuration_never_initializes_llm_client():
    saved = (
        aqv.RUN_MODE,
        aqv.PURE_REF_PREFILTER,
        aqv.DETERMINISTIC_SOURCE_SPLITTER,
        aqv.FREE_NO_LLM,
        aqv.REF_DISAMBIG_FALLBACK,
        aqv.SUPRA_LINKING_AGGRESSIVENESS,
        aqv.client,
        aqv.LINK_RESOLVER,
    )
    args = Namespace(run_mode="free", dry_fire=True, no_a2aj=True)
    try:
        with patch.object(aqv, "_ensure_llm_client") as ensure:
            aqv._configure_from_args(args)
        ensure.assert_not_called()
        assert aqv.PURE_REF_PREFILTER
        assert aqv.DETERMINISTIC_SOURCE_SPLITTER
        assert aqv.FREE_NO_LLM
        assert not aqv.REF_DISAMBIG_FALLBACK
    finally:
        (
            aqv.RUN_MODE,
            aqv.PURE_REF_PREFILTER,
            aqv.DETERMINISTIC_SOURCE_SPLITTER,
            aqv.FREE_NO_LLM,
            aqv.REF_DISAMBIG_FALLBACK,
            aqv.SUPRA_LINKING_AGGRESSIVENESS,
            aqv.client,
            aqv.LINK_RESOLVER,
        ) = saved


def test_local_only_forces_free_mode_and_disables_external_sources():
    saved = (
        aqv.LOCAL_ONLY, aqv.RUN_MODE, aqv.FREE_NO_LLM,
        aqv.USE_A2AJ, aqv.USE_DB_SEARCH, aqv.LINK_RESOLVER,
    )

    class CitationDB:
        external_enabled = True
        def set_external_enabled(self, enabled):
            self.external_enabled = enabled

    citation_db = CitationDB()
    try:
        with patch.object(aqv, "_ensure_llm_client") as ensure, \
                patch.object(
                    aqv._provider_registry, "get_citation_db",
                    return_value=citation_db,
                ), patch.object(
                    aqv._provider_registry, "build_" + "resolver", create=True,
                ) as external_builder, patch.object(
                    aqv.a2aj_client, "set_local_only",
                ) as set_local:
            aqv._configure_from_args(Namespace(
                run_mode="high_accuracy", local_only=True, dry_fire=False,
                use_a2aj=True, use_db_search=True,
            ))

        ensure.assert_not_called()
        external_builder.assert_not_called()
        set_local.assert_called_once_with(True)
        assert aqv.LOCAL_ONLY and aqv.RUN_MODE == "free" and aqv.FREE_NO_LLM
        assert aqv.USE_A2AJ and aqv.USE_DB_SEARCH
        assert citation_db.external_enabled is False
        assert aqv.LINK_RESOLVER.resolve_url(" https://example.test ") == (
            "https://example.test"
        )
    finally:
        (
            aqv.LOCAL_ONLY, aqv.RUN_MODE, aqv.FREE_NO_LLM,
            aqv.USE_A2AJ, aqv.USE_DB_SEARCH, aqv.LINK_RESOLVER,
        ) = saved


def test_economy_keeps_original_pure_reference_scope():
    saved = (
        aqv.RUN_MODE,
        aqv.PURE_REF_PREFILTER,
        aqv.DETERMINISTIC_SOURCE_SPLITTER,
        aqv.FREE_NO_LLM,
        aqv.REF_DISAMBIG_FALLBACK,
        aqv.SUPRA_LINKING_AGGRESSIVENESS,
        aqv.client,
        aqv.LINK_RESOLVER,
    )
    try:
        with patch.object(aqv, "_ensure_llm_client", return_value=None):
            aqv._configure_from_args(Namespace(
                run_mode="economy", dry_fire=True, no_a2aj=True
            ))
        assert aqv.PURE_REF_PREFILTER
        assert not aqv.DETERMINISTIC_SOURCE_SPLITTER
        assert not aqv.FREE_NO_LLM
        assert aqv.REF_DISAMBIG_FALLBACK

        with patch.object(aqv, "_ensure_llm_client", return_value=None):
            aqv._configure_from_args(Namespace(
                run_mode="ultra_economy", dry_fire=True, no_a2aj=True
            ))
        assert aqv.PURE_REF_PREFILTER
        assert aqv.DETERMINISTIC_SOURCE_SPLITTER
        assert not aqv.FREE_NO_LLM
        assert aqv.REF_DISAMBIG_FALLBACK
    finally:
        (
            aqv.RUN_MODE,
            aqv.PURE_REF_PREFILTER,
            aqv.DETERMINISTIC_SOURCE_SPLITTER,
            aqv.FREE_NO_LLM,
            aqv.REF_DISAMBIG_FALLBACK,
            aqv.SUPRA_LINKING_AGGRESSIVENESS,
            aqv.client,
            aqv.LINK_RESOLVER,
        ) = saved


def test_supra_linking_setting_is_independent_of_run_mode():
    saved = aqv.SUPRA_LINKING_AGGRESSIVENESS
    try:
        with patch.object(aqv, "_ensure_llm_client", return_value=None):
            aqv._configure_from_args(Namespace(
                run_mode="high_accuracy", supra_linking="aggressive",
                dry_fire=True, use_a2aj=False,
            ))
        assert aqv.SUPRA_LINKING_AGGRESSIVENESS == "aggressive"
    finally:
        aqv.SUPRA_LINKING_AGGRESSIVENESS = saved
